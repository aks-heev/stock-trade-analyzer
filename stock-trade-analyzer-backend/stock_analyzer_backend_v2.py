from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Dict, Optional
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
import json
from datetime import datetime
import io
import traceback

app = FastAPI(title="Stock Trade Analyzer API", version="2.0.0")

# Enable CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Pydantic Models
class TradeData(BaseModel):
    symbol: str
    entry_price: float
    entry_date: str
    exit_price: Optional[float] = None
    exit_date: Optional[str] = None
    quantity: int
    pnl: float
    pnl_percent: float
    status: str  # "Profitable", "Loss", "Open"
    commission: float = 0

class PositionData(BaseModel):
    symbol: str
    quantity: int
    avg_price: float
    total_cost: float

class AnalysisResult(BaseModel):
    trades: List[TradeData]
    positions: List[PositionData]
    metrics: Dict
    summary: Dict

# Helper Functions
def detect_data_start_row(file_content: bytes) -> int:
    """
    Auto-detect where actual data starts by looking for required column headers.
    Returns the row number to skip to, or 0 if data starts at row 0.
    """
    try:
        # Read without skipping to examine headers
        excel_file = io.BytesIO(file_content)
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active
        
        required_keywords = ['symbol', 'isin', 'scrip', 'date', 'order', 'price', 'quantity', 'qty']
        
        # Check first 10 rows for header row
        for row_idx in range(1, min(11, ws.max_row + 1)):
            row_values = [str(cell.value).lower() if cell.value else '' for cell in ws[row_idx]]
            row_text = ' '.join(row_values)
            
            # Count how many required keywords are in this row
            match_count = sum(1 for keyword in required_keywords if keyword in row_text)
            
            # If this row has multiple keywords, it's likely the header
            if match_count >= 3:
                return row_idx  # Return the row number (1-indexed)
        
        return 0  # Default: no headers found, assume data starts at row 1
    
    except Exception as e:
        print(f"Auto-detect warning: {e}, using default skip_rows=5")
        return 5


def parse_excel_file(file_content: bytes, skip_rows: Optional[int] = None) -> tuple:
    """
    Parse Excel file and return DataFrame with actual skip_rows used.
    
    Args:
        file_content: Raw file bytes
        skip_rows: Number of rows to skip (None = auto-detect, default tries 5 first)
    
    Returns:
        Tuple of (DataFrame, rows_skipped)
    """
    try:
        excel_file = io.BytesIO(file_content)
        
        # Determine skip_rows strategy
        if skip_rows is not None:
            # User specified skip_rows
            df = pd.read_excel(excel_file, skiprows=skip_rows)
            rows_skipped = skip_rows
        else:
            # Try default of 5 first (common for broker exports)
            try:
                df = pd.read_excel(excel_file, skiprows=5)
                rows_skipped = 5
                
                # Verify we got actual data with expected columns
                df_normalized = df.copy()
                df_normalized.columns = df_normalized.columns.str.strip().str.lower()
                
                expected_cols = ['symbol', 'date', 'order_type', 'price', 'quantity']
                expected_present = sum(1 for col in expected_cols if col in df_normalized.columns or any(alt in df_normalized.columns for alt in ['isin', 'scrip', 'side', 'buy_sell', 'qty', 'shares', 'execution date and time', 'type', 'value']))
                
                if expected_present < 3:
                    # Not a good match, try auto-detect
                    raise ValueError("Expected columns not found with skip_rows=5")
            
            except Exception:
                # Fall back to auto-detection
                excel_file.seek(0)
                detected_skip = detect_data_start_row(excel_file.getvalue())
                excel_file.seek(0)
                
                if detected_skip > 0:
                    df = pd.read_excel(excel_file, skiprows=detected_skip - 1)
                    rows_skipped = detected_skip - 1
                else:
                    # Last resort: read without skipping
                    excel_file.seek(0)
                    df = pd.read_excel(excel_file)
                    rows_skipped = 0
        
        if df.empty:
            raise ValueError("Resulting dataframe is empty")
        
        return df, rows_skipped
    
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error reading Excel file: {str(e)}")


def normalize_column_names(df: pd.DataFrame) -> pd.DataFrame:
    """Normalize column names to standard format"""
    df.columns = df.columns.str.strip().str.lower()
    return df


def match_trades_fifo(orders: pd.DataFrame) -> tuple:
    """
    Match BUY and SELL orders using FIFO (First In, First Out) algorithm
    Tracks average price for each stock
    Returns: (trades_list, positions_list)
    """
    print(f"DEBUG match_trades_fifo: Starting with {len(orders)} rows")
    print(f"DEBUG match_trades_fifo: Columns available: {list(orders.columns)}")
    
    trades = []
    open_positions = {}
    stock_avg_prices = {}  # Track average price per stock
    
    # Sort by date
    try:
        print(f"DEBUG match_trades_fifo: Attempting to sort by 'date'")
        print(f"DEBUG match_trades_fifo: Date column type: {orders['date'].dtype}")
        print(f"DEBUG match_trades_fifo: Sample date values: {orders['date'].head(3).tolist()}")
        orders = orders.sort_values('date')
        print(f"DEBUG match_trades_fifo: Sort successful")
    except KeyError as ke:
        print(f"ERROR in match_trades_fifo: KeyError when accessing 'date' - {ke}")
        print(f"ERROR: Available columns: {list(orders.columns)}")
        raise
    except Exception as e:
        print(f"ERROR in match_trades_fifo: {type(e).__name__}: {e}")
        raise
    
    for idx, (_, row) in enumerate(orders.iterrows()):
        try:
            symbol = str(row.get('symbol', 'UNKNOWN')).upper().strip()
            order_type = str(row.get('order_type', 'BUY')).upper().strip()
            price = float(row.get('price', 0)) if pd.notna(row.get('price')) else 0
            quantity = int(float(row.get('quantity', 0))) if pd.notna(row.get('quantity')) else 0
            date = row['date']
            commission = float(row.get('commission', 0)) if pd.notna(row.get('commission', 0)) else 0
            
            if quantity <= 0:
                continue
            
            if symbol not in open_positions:
                open_positions[symbol] = {'buys': [], 'sells': []}
                stock_avg_prices[symbol] = {'quantity': 0, 'total_cost': 0, 'avg_price': 0}
            
            if 'BUY' in order_type:
                # Add to open positions
                open_positions[symbol]['buys'].append({
                    'price': price,
                    'quantity': quantity,
                    'date': date,
                    'commission': commission
                })
                
                # Update average price for this stock
                current_qty = stock_avg_prices[symbol]['quantity']
                current_cost = stock_avg_prices[symbol]['total_cost']
                
                new_total_cost = current_cost + (price * quantity) + commission
                new_total_qty = current_qty + quantity
                
                stock_avg_prices[symbol]['quantity'] = new_total_qty
                stock_avg_prices[symbol]['total_cost'] = new_total_cost
                stock_avg_prices[symbol]['avg_price'] = new_total_cost / new_total_qty if new_total_qty > 0 else 0
                
                print(f"  BUY {symbol}: {quantity} @ ₹{price:.2f} | Avg Price: ₹{stock_avg_prices[symbol]['avg_price']:.2f} | Total Qty: {new_total_qty}")
                
            elif 'SELL' in order_type:
                # Match with oldest BUY orders (FIFO)
                remaining_qty = quantity
                
                while remaining_qty > 0 and open_positions[symbol]['buys']:
                    buy = open_positions[symbol]['buys'][0]
                    
                    if buy['quantity'] <= remaining_qty:
                        # Complete match
                        matched_qty = buy['quantity']
                        pnl = (price - buy['price']) * matched_qty - (buy['commission'] + commission)
                        pnl_percent = ((price - buy['price']) / buy['price'] * 100) if buy['price'] != 0 else 0
                        
                        trades.append({
                            'symbol': symbol,
                            'entry_price': buy['price'],
                            'entry_date': str(buy['date']),
                            'exit_price': price,
                            'exit_date': str(date),
                            'quantity': matched_qty,
                            'pnl': round(pnl, 2),
                            'pnl_percent': round(pnl_percent, 2),
                            'status': 'Profitable' if pnl > 0 else 'Loss' if pnl < 0 else 'Breakeven',
                            'commission': round(buy['commission'] + commission, 2)
                        })
                        
                        remaining_qty -= matched_qty
                        open_positions[symbol]['buys'].pop(0)
                    else:
                        # Partial match
                        matched_qty = remaining_qty
                        pnl = (price - buy['price']) * matched_qty - (buy['commission'] + commission)
                        pnl_percent = ((price - buy['price']) / buy['price'] * 100) if buy['price'] != 0 else 0
                        
                        trades.append({
                            'symbol': symbol,
                            'entry_price': buy['price'],
                            'entry_date': str(buy['date']),
                            'exit_price': price,
                            'exit_date': str(date),
                            'quantity': matched_qty,
                            'pnl': round(pnl, 2),
                            'pnl_percent': round(pnl_percent, 2),
                            'status': 'Profitable' if pnl > 0 else 'Loss' if pnl < 0 else 'Breakeven',
                            'commission': round(buy['commission'] + commission, 2)
                        })
                        
                        buy['quantity'] -= matched_qty
                        remaining_qty = 0
                
                # Update average price after sell (reduce quantity)
                current_qty = stock_avg_prices[symbol]['quantity']
                current_cost = stock_avg_prices[symbol]['total_cost']
                
                # Reduce by sold quantity proportionally
                if current_qty > 0:
                    cost_per_unit = current_cost / current_qty
                    reduction = quantity * cost_per_unit
                    
                    new_total_qty = max(0, current_qty - quantity)
                    new_total_cost = max(0, current_cost - reduction)
                    
                    stock_avg_prices[symbol]['quantity'] = new_total_qty
                    stock_avg_prices[symbol]['total_cost'] = new_total_cost
                    stock_avg_prices[symbol]['avg_price'] = new_total_cost / new_total_qty if new_total_qty > 0 else 0
                    
                    print(f"  SELL {symbol}: {quantity} @ ₹{price:.2f} | Remaining Qty: {new_total_qty} | Avg Price: ₹{stock_avg_prices[symbol]['avg_price']:.2f}")
                
        except Exception as e:
            print(f"ERROR in match_trades_fifo row {idx}: {type(e).__name__} - {e}")
            raise
    
    # Add open positions
    for symbol, positions in open_positions.items():
        for buy in positions['buys']:
            trades.append({
                'symbol': symbol,
                'entry_price': buy['price'],
                'entry_date': str(buy['date']),
                'exit_price': None,
                'exit_date': None,
                'quantity': buy['quantity'],
                'pnl': 0,
                'pnl_percent': 0,
                'status': 'Open',
                'commission': buy['commission']
            })
    
    # Build positions list from stock_avg_prices
    print(f"\nDEBUG match_trades_fifo: Final Average Prices by Stock:")
    positions = []
    for symbol, avg_data in stock_avg_prices.items():
        if avg_data['quantity'] > 0:
            positions.append({
                'symbol': symbol,
                'quantity': avg_data['quantity'],
                'avg_price': round(avg_data['avg_price'], 2),
                'total_cost': round(avg_data['total_cost'], 2)
            })
            print(f"  {symbol}: Avg Price = ₹{avg_data['avg_price']:.2f} | Remaining Qty: {avg_data['quantity']} | Total Cost: ₹{avg_data['total_cost']:.2f}")
    
    print(f"DEBUG match_trades_fifo: Created {len(trades)} trades, {len(positions)} open positions\n")
    return trades, positions


def calculate_metrics(trades: List[Dict]) -> Dict:
    """Calculate portfolio metrics from matched trades"""
    if not trades:
        return {}
    
    closed_trades = [t for t in trades if t['status'] != 'Open']
    open_trades = [t for t in trades if t['status'] == 'Open']
    
    if not closed_trades:
        return {
            'total_trades': len(trades),
            'closed_trades': 0,
            'open_trades': len(open_trades),
            'win_rate': 0,
            'total_pnl': 0,
            'avg_win': 0,
            'avg_loss': 0,
            'best_trade': 0,
            'worst_trade': 0,
            'profit_factor': 0
        }
    
    winning_trades = [t for t in closed_trades if t['pnl'] > 0]
    losing_trades = [t for t in closed_trades if t['pnl'] < 0]
    
    total_pnl = sum(t['pnl'] for t in closed_trades)
    win_rate = (len(winning_trades) / len(closed_trades) * 100) if closed_trades else 0
    
    avg_win = sum(t['pnl'] for t in winning_trades) / len(winning_trades) if winning_trades else 0
    avg_loss = abs(sum(t['pnl'] for t in losing_trades) / len(losing_trades)) if losing_trades else 0
    
    best_trade = max((t['pnl'] for t in closed_trades), default=0)
    worst_trade = min((t['pnl'] for t in closed_trades), default=0)
    
    gross_profit = sum(t['pnl'] for t in winning_trades) if winning_trades else 0
    gross_loss = abs(sum(t['pnl'] for t in losing_trades)) if losing_trades else 0
    
    profit_factor = gross_profit / gross_loss if gross_loss > 0 else 0
    
    return {
        'total_trades': len(trades),
        'closed_trades': len(closed_trades),
        'open_trades': len(open_trades),
        'winning_trades': len(winning_trades),
        'losing_trades': len(losing_trades),
        'win_rate': round(win_rate, 2),
        'total_pnl': round(total_pnl, 2),
        'avg_win': round(avg_win, 2),
        'avg_loss': round(avg_loss, 2),
        'best_trade': round(best_trade, 2),
        'worst_trade': round(worst_trade, 2),
        'profit_factor': round(profit_factor, 2)
    }


# API Routes
@app.get("/")
@app.head("/")
async def root():
    return {
        "message": "Stock Trade Analyzer API v2.0.0",
        "endpoints": [
            "/analyze-xlsx - POST: Upload Excel file for analysis (query param: skip_rows=5)",
            "/health - GET: Check API health"
        ]
    }


@app.get("/health")
async def health_check():
    return {"status": "healthy", "timestamp": datetime.now().isoformat()}


@app.post("/analyze-xlsx", response_model=AnalysisResult)
async def analyze_xlsx(file: UploadFile = File(...), skip_rows: Optional[int] = None):
    """
    Analyze uploaded Excel file with trade data.
    
    Handles files with broker header information by:
    1. Trying skip_rows parameter if provided
    2. Defaulting to skip_rows=5 (common broker format)
    3. Auto-detecting header row if default doesn't work
    
    Args:
        file: Excel file (.xlsx)
        skip_rows: Optional - number of header rows to skip (default: auto-detect, usually 5)
    
    Expected columns: Symbol, Date, Order Type, Price, Quantity, Commission
    """
    try:
        print("\n" + "="*80)
        print("ANALYZING FILE")
        print("="*80)
        
        # Read file with smart skip_rows handling
        contents = await file.read()
        print(f"File received: {file.filename}")
        
        df, rows_skipped = parse_excel_file(contents, skip_rows=skip_rows)
        print(f"✓ Parsed Excel file with skiprows={rows_skipped}")
        print(f"  Columns before normalization: {list(df.columns)}")
        
        # Normalize columns
        df = normalize_column_names(df)
        print(f"✓ Normalized column names")
        print(f"  Columns after normalization: {list(df.columns)}")
        
        # Map alternative column names to standard format
        # This handles broker-specific column naming conventions
        column_mapping = {
            'symbol': ['symbol', 'isin', 'scrip', 'stock symbol', 'ticker', 'stock name'],
            'date': ['date', 'date_time', 'order_date', 'execution date and time', 'execution date', 'trade date'],
            'order_type': ['order_type', 'side', 'buy_sell', 'bs', 'type', 'order status', 'trans type'],
            'price': ['price', 'rate', 'bid', 'ask', 'value', 'trade price', 'execution price'],
            'quantity': ['quantity', 'qty', 'shares', 'no of units'],
            'commission': ['commission', 'brokerage', 'fees', 'charges', 'transaction charges']
        }
        
        # Apply column mapping
        print(f"✓ Applying column mapping...")
        price_from_value = False
        
        # Check if 'value' column exists before mapping
        if 'value' in df.columns:
            price_from_value = True
            print(f"  ⚠ Detected 'Value' column - will need to divide by quantity")
        
        for standard_col, alternatives in column_mapping.items():
            if standard_col not in df.columns:
                for alt in alternatives:
                    if alt in df.columns:
                        df[standard_col] = df[alt]
                        print(f"  Mapped '{alt}' → '{standard_col}'")
                        break
        
        print(f"✓ Columns after mapping: {list(df.columns)}")
        
        # Validate we have minimum required columns
        required_columns = ['symbol', 'date', 'order_type', 'price', 'quantity']
        missing_cols = [col for col in required_columns if col not in df.columns]
        
        if missing_cols:
            print(f"✗ Missing required columns: {missing_cols}")
            print(f"  Available: {list(df.columns)}")
            raise ValueError(f"Missing required columns: {missing_cols}. Available columns: {list(df.columns)}")
        
        print(f"✓ All required columns present")
        
        # Parse dates FIRST - handle multiple date formats
        if 'date' in df.columns:
            print(f"✓ Processing date column...")
            print(f"  Sample values: {df['date'].head(3).tolist()}")
            print(f"  Data type: {df['date'].dtype}")
            
            try:
                # Try the specific broker date format first
                df['date'] = pd.to_datetime(df['date'], format='%d-%m-%Y %H:%M %p', errors='coerce')
                # For rows that didn't parse, try generic parsing
                mask = df['date'].isna()
                if mask.any():
                    print(f"  ⚠ {mask.sum()} rows failed specific format, trying generic...")
                    df.loc[mask, 'date'] = pd.to_datetime(df.loc[mask, 'date'], errors='coerce')
            except Exception as e:
                print(f"  Exception during date parsing: {e}")
                df['date'] = pd.to_datetime(df['date'], errors='coerce')
            
            print(f"  After parsing - type: {df['date'].dtype}")
            print(f"  Sample values: {df['date'].head(3).tolist()}")
        
        # Remove rows with NaT (invalid) dates before data type conversions
        initial_rows = len(df)
        df = df.dropna(subset=['date'])
        print(f"✓ Dropped {initial_rows - len(df)} rows with invalid dates")
        
        # Clean and convert data types
        print(f"✓ Converting data types...")
        if 'quantity' in df.columns:
            df['quantity'] = pd.to_numeric(df['quantity'], errors='coerce').fillna(0).astype(int)
            print(f"  Quantity: {df['quantity'].dtype} (min={df['quantity'].min()}, max={df['quantity'].max()})")
        
        if 'price' in df.columns:
            df['price'] = pd.to_numeric(df['price'], errors='coerce').fillna(0)
            # If price came from 'value' column, divide by quantity to get unit price
            if price_from_value:
                print(f"  ⚠ Price came from 'Value' column, calculating unit price = Value / Quantity")
                df['price'] = df['price'] / df['quantity'].replace(0, 1)  # Avoid division by zero
            print(f"  Price: {df['price'].dtype} (min={df['price'].min():.2f}, max={df['price'].max():.2f})")
        
        if 'commission' in df.columns:
            df['commission'] = pd.to_numeric(df['commission'], errors='coerce').fillna(0)
        else:
            df['commission'] = 0.0
        print(f"  Commission added")
        
        # Clean order_type - extract BUY or SELL
        if 'order_type' in df.columns:
            print(f"✓ Processing order_type column...")
            print(f"  Unique values: {df['order_type'].unique()}")
            df['order_type'] = df['order_type'].astype(str).str.upper().str.strip()
            initial_len = len(df)
            df = df[df['order_type'].isin(['BUY', 'SELL'])]
            print(f"  Filtered from {initial_len} to {len(df)} rows (kept only BUY/SELL)")
        
        # Remove rows with missing critical data
        initial_len = len(df)
        df = df.dropna(subset=['symbol', 'date', 'order_type', 'price', 'quantity'])
        print(f"✓ Dropped {initial_len - len(df)} rows with missing critical data")
        
        df = df[df['quantity'] > 0]  # Only positive quantities
        
        print(f"✓ Final dataset: {len(df)} rows, {len(df.columns)} columns")
        
        if df.empty:
            raise ValueError("No valid trade data found after cleaning")
        
        # Match trades
        print(f"✓ Matching trades with FIFO algorithm...")
        trades, positions = match_trades_fifo(df)
        print(f"✓ Matched {len(trades)} trades, {len(positions)} open positions")
        
        # Calculate metrics
        metrics = calculate_metrics(trades)
        print(f"✓ Calculated metrics")
        
        # Prepare summary
        summary = {
            'file_name': file.filename,
            'rows_processed': len(df),
            'rows_skipped': rows_skipped,
            'analysis_date': datetime.now().isoformat(),
            'total_symbols': len(set(t['symbol'] for t in trades))
        }
        
        print(f"✓ Analysis complete!")
        print("="*80 + "\n")
        
        return AnalysisResult(
            trades=trades,
            positions=positions,
            metrics=metrics,
            summary=summary
        )
    
    except Exception as e:
        print(f"\n✗ ERROR: {type(e).__name__}: {str(e)}")
        print("="*80)
        print(traceback.format_exc())
        print("="*80 + "\n")
        raise HTTPException(
            status_code=400,
            detail=f"Error processing file: {str(e)}"
        )


@app.post("/single-trade")
async def analyze_single_trade(data: Dict):
    """
    Analyze a single trade
    Expected: {
        "entry_price": float,
        "exit_price": float,
        "quantity": int,
        "commission": float
    }
    """
    try:
        entry_price = float(data.get('entry_price', 0))
        exit_price = float(data.get('exit_price', 0))
        quantity = int(data.get('quantity', 1))
        commission = float(data.get('commission', 0))
        
        pnl = (exit_price - entry_price) * quantity - commission
        pnl_percent = ((exit_price - entry_price) / entry_price * 100) if entry_price != 0 else 0
        
        return {
            'entry_price': entry_price,
            'exit_price': exit_price,
            'quantity': quantity,
            'pnl': round(pnl, 2),
            'pnl_percent': round(pnl_percent, 2),
            'status': 'Profitable' if pnl > 0 else 'Loss' if pnl < 0 else 'Breakeven',
            'commission': commission
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.get("/upstox/callback")
async def upstox_callback(code: str | None = None, state: str | None = None, request: Request = None):
    # 1) Validate state (if you use it)
    # 2) Exchange `code` for access token using Upstox token API
    # 3) Store tokens in DB/file and redirect user to a success page
    return {"message": "Upstox callback received", "code": code, "state": state}


if __name__ == "__main__":
    import uvicorn
    import os
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)